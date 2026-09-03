import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from extract import df_stock_days
from transform import transform_stock_days

def export_excel(df, ruta_salida):
    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)

    df.to_excel(ruta_salida, index=False, sheet_name="SKU_Clasificados")

    wb = openpyxl.load_workbook(ruta_salida)
    ws = wb["SKU_Clasificados"]

    fill_critico = PatternFill(
        start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
    )

    fill_alerta = PatternFill(
        start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
    )

    fill_normal = PatternFill(
        start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
    )

    font_bold = Font(bold=True)

    encabezados = [celda.value for celda in ws[1]]
    col_idx = (
        encabezados.index("categoria_inventario") + 1
    )

    for fila in range(2, ws.max_row + 1):
        celda_categoria = ws.cell(row=fila, column=col_idx)
        valor = celda_categoria.value

        if valor == "Crítico":
            celda_categoria.fill = fill_critico
            celda_categoria.font = font_bold
        elif valor == "Alerta":
            celda_categoria.fill = fill_alerta
            celda_categoria.font = font_bold
        elif valor == "Normal":
            celda_categoria.fill = fill_normal

    wb.save(ruta_salida)
    print(f"✅ Excel exportado y formateado con éxito en: {ruta_salida}")

if __name__ == "__main__":
        df_clasificado = transform_stock_days(df_stock_days)
        ruta_archivo = os.path.join("data", "reporte_clasificacion_sku.xlsx")
        export_excel(df_clasificado, ruta_archivo)
        print(f"✅ Excel exportado y formateado con éxito en: {ruta_archivo}")